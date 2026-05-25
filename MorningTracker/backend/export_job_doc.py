import sys
import os
from datetime import datetime
import io

# Add current directory to path so we can import db
sys.path.append(os.path.dirname(os.path.abspath(__file__)))

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Error: python-docx is not installed. Please install it with 'pip install python-docx'")
    sys.exit(1)

from db.database import get_db_sync, Article, ScrapeJob
from sqlalchemy import select

def export_job_docx(job_id):
    print(f"Exporting data for job: {job_id}")
    
    with get_db_sync() as db:
        # Get job
        job = db.query(ScrapeJob).filter(ScrapeJob.id == job_id).first()
        if not job:
            # Try partial match if not found (sometimes user gives short ID)
            job = db.query(ScrapeJob).filter(ScrapeJob.id.like(f"{job_id}%")).first()
            if not job:
                print(f"Job {job_id} not found in database.")
                return None
            else:
                print(f"Found job: {job.id} for input: {job_id}")
                job_id = job.id

        # Get articles
        articles = db.query(Article).filter(Article.scrape_job_id == job_id).order_by(Article.published_at.desc()).all()
        print(f"Found {len(articles)} articles.")

        doc = Document()
        
        # Header Section
        title = doc.add_heading(f"NEXUS News Report: {job.sector}", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        metadata = doc.add_paragraph()
        metadata.add_run(f"Report Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        metadata.add_run(f"Sector: {job.sector} | Region: {job.region.upper()}\n")
        metadata.add_run(f"Articles Found: {len(articles)}")
        metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_page_break()

        for a in articles:
            # Article Title
            h = doc.add_heading(a.title, level=1)
            
            # Article Metadata
            p = doc.add_paragraph()
            p.add_run(f"Source: {a.agency or 'Unknown'} | Date: {a.published_at.strftime('%Y-%m-%d') if a.published_at else 'N/A'}\n").italic = True
            url_text = a.resolved_url or a.url
            p.add_run(f"URL: {url_text}").font.color.rgb = RGBColor(0, 0, 255)
            
            # Summary Section
            if a.summary:
                doc.add_heading("Summary", level=2)
                doc.add_paragraph(a.summary)
            
            # Full Content Section
            if a.full_body:
                doc.add_heading("Full Article Content", level=2)
                doc.add_paragraph(a.full_body)
            
            doc.add_paragraph("_" * 50) # Separator

        brand_name = job.sector.replace(" ", "_")
        filename = f"NEXUS_Report_{brand_name}_{job_id}_{datetime.now().strftime('%Y-%m-%d')}.docx"
        output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", filename)
        
        doc.save(output_path)
        print(f"Successfully saved report to: {output_path}")
        return output_path

if __name__ == "__main__":
    job_id = "70b4612f"
    if len(sys.argv) > 1:
        job_id = sys.argv[1]
    
    path = export_job_docx(job_id)
    if path:
        print(f"\nRESULT_FILE_PATH: {path}")

from fastapi import APIRouter, Query, HTTPException, WebSocket, WebSocketDisconnect, Depends, status
import asyncio
import csv
import io
from typing import Optional, AsyncGenerator
from datetime import date, datetime
from sqlalchemy import select, func, or_, and_, update, desc, text, delete
from db.database import get_db, Article, ScrapeJob
from .auth_utils import get_auth_user as get_current_user, TokenData
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

router = APIRouter()

@router.get("/")
async def get_articles(
    sector: Optional[str] = None,
    region: Optional[str] = None,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    job_id: Optional[str] = None,
    search: Optional[str] = None,
    has_body: Optional[bool] = None,
    page: int = 1,
    page_size: int = 25,
    current_user: TokenData = Depends(get_current_user)
):
    """Fetch articles with filtering and pagination (Global for admin)."""
    offset = (page - 1) * page_size
    async with get_db() as db:
        stmt = select(Article)
        if not current_user.is_admin:
            stmt = stmt.where(Article.user_id == current_user.id)
        
        if sector: stmt = stmt.where(Article.sector == sector)
        if region: stmt = stmt.where(Article.region == region)
        if date_from: stmt = stmt.where(Article.published_at >= date_from)
        if date_to: stmt = stmt.where(Article.published_at <= date_to)
        if job_id: stmt = stmt.where(Article.scrape_job_id == job_id)
        if search:
            stmt = stmt.where(or_(
                Article.title.ilike(f"%{search}%"),
                Article.full_body.ilike(f"%{search}%")
            ))
        if has_body is True:
            stmt = stmt.where(and_(Article.full_body != None, func.length(Article.full_body) > 100))
        elif has_body is False:
            stmt = stmt.where(or_(Article.full_body == None, func.length(Article.full_body) <= 100))

        # Count total
        count_stmt = select(func.count()).select_from(stmt.subquery())
        res_total = await db.execute(count_stmt)
        total = res_total.scalar()

        # Get results
        stmt = stmt.order_by(Article.published_at.desc()).offset(offset).limit(page_size)
        res_articles = await db.execute(stmt)
        articles = res_articles.scalars().all()

    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_pages": -(-total // page_size) if total else 0,
        "articles": articles
    }

async def _fetch_stats_logic(user_id: str, is_admin: bool = False):
    """Internal stats logic (Global for admins)."""
    async with get_db() as db:
        # Build base filters
        def apply_user_filter(query):
            if not is_admin:
                return query.where(Article.user_id == user_id)
            return query

        def apply_job_user_filter(query):
            if not is_admin:
                return query.where(ScrapeJob.user_id == user_id)
            return query

        # Total
        total_res = await db.execute(apply_user_filter(select(func.count(Article.id))))
        total = total_res.scalar() or 0
        
        # With Body
        body_res = await db.execute(apply_user_filter(select(func.count(Article.id)).where(Article.full_body != None, func.length(Article.full_body) > 100)))
        with_body = body_res.scalar() or 0
        
        # With Summary
        sum_res = await db.execute(apply_user_filter(select(func.count(Article.id)).where(Article.summary != None)))
        with_summary = sum_res.scalar() or 0
        
        # By Sector
        sect_res = await db.execute(apply_user_filter(select(Article.sector, func.count(Article.id)).group_by(Article.sector).order_by(desc(func.count(Article.id)))))
        by_sector = [{"sector": r[0], "count": r[1]} for r in sect_res.all()]
        
        # Jobs
        jobs_res = await db.execute(apply_job_user_filter(select(ScrapeJob.status, func.count(ScrapeJob.id)).group_by(ScrapeJob.status)))
        jobs_by_status = [{"status": r[0], "count": r[1]} for r in jobs_res.all()]

    return {
        "total_articles": total,
        "articles_with_body": with_body,
        "articles_with_summary": with_summary,
        "body_coverage_pct": round((with_body / total * 100), 1) if total else 0,
        "by_sector": by_sector,
        "jobs_by_status": jobs_by_status
    }

@router.get("/stats/summary")
async def get_stats(current_user: TokenData = Depends(get_current_user)):
    return await _fetch_stats_logic(current_user.id, current_user.is_admin)

@router.get("/export/csv")
async def export_csv(
    job_id: str,
    current_user: TokenData = Depends(get_current_user)
):
    async def generate():
        yield b'\xef\xbb\xbf' # BOM for Excel
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Title", "URL", "Agency", "Published_At", "Summary", "Full Body"])
        yield output.getvalue().encode("utf-8")
        
        async with get_db() as db:
            stmt = select(Article).where(Article.scrape_job_id == job_id)
            if not current_user.is_admin:
                stmt = stmt.where(Article.user_id == current_user.id)
                
            res = await db.execute(stmt)
            for a in res.scalars():
                out = io.StringIO()
                cw = csv.writer(out)
                cw.writerow([a.title, a.url, a.agency, a.published_at, a.summary, a.full_body])
                yield out.getvalue().encode("utf-8")

    return StreamingResponse(generate(), media_type="text/csv", headers={"Content-Disposition": f"attachment; filename=export_{job_id}.csv"})

@router.get("/export/xlsx")
async def export_xlsx(
    job_id: str,
    current_user: TokenData = Depends(get_current_user)
):
    """Excel export with formatting (Blueprint Phase 5)."""
    async with get_db() as db:
        # Get job for naming
        stmt_job = select(ScrapeJob).where(ScrapeJob.id == job_id)
        if not current_user.is_admin:
            stmt_job = stmt_job.where(ScrapeJob.user_id == current_user.id)
            
        job_res = await db.execute(stmt_job)
        job = job_res.scalar_one_or_none()
        if not job:
            raise HTTPException(404, "Job not found or access denied")
        
        # Get articles
        stmt_count = select(func.count()).where(Article.scrape_job_id == job_id)
        if not current_user.is_admin:
            stmt_count = stmt_count.where(Article.user_id == current_user.id)
            
        total_count = (await db.execute(stmt_count)).scalar() or 0
        
        if total_count == 0:
            raise HTTPException(400, "No articles found for this job. Ensure discovery is complete.")
            
        if total_count > 5000:
            # Hard limit for XLSX to prevent OOM
            raise HTTPException(400, "Job too large for XLSX (Max 5,000 articles). Use CSV export instead.")

        stmt_articles = select(Article).where(Article.scrape_job_id == job_id)
        if not current_user.is_admin:
            stmt_articles = stmt_articles.where(Article.user_id == current_user.id)
            
        stmt_articles = stmt_articles.order_by(Article.published_at.desc())
        res = await db.execute(stmt_articles)
        articles = res.scalars().all()

        wb = Workbook()
        ws = wb.active
        brand_name = job.sector.replace(" ", "_")
        ws.title = f"Nexus_{brand_name}"

        # Columns: Title, Resolved URL, Publisher/Agency, Author, Summary, Full Body, Published At, Source Feed, Keyword Matched
        headers = ["Title", "Resolved URL", "Publisher/Agency", "Author", "Summary", "Full Body", "Published At", "Source Feed", "Keyword Matched"]
        ws.append(headers)

        # Formatting Header
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Freeze top row
        ws.freeze_panes = "A2"

        # Data Rows
        alternate_fill = PatternFill(start_color="F0F4FA", end_color="F0F4FA", fill_type="solid")
        for i, a in enumerate(articles, start=2):
            published_str = a.published_at.strftime("%d %b %Y %H:%M") if a.published_at else ""
            
            # Robust Author Selection
            primary_author = a.author
            if not primary_author and a.extra_metadata:
                # Try fallback to stored metadata if enrichment didn't run yet
                meta = a.extra_metadata.get("author_metadata", {})
                if isinstance(meta, dict):
                    primary_author = meta.get("name")
            
            row_data = [
                a.title,
                a.resolved_url or a.url,
                a.agency or "Unknown Publisher",
                primary_author or "Staff Reporter",
                a.summary,
                a.full_body,
                published_str,
                a.source_feed or "google_news",
                job.sector
            ]
            ws.append(row_data)
            
            # Formatting for the new row
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=i, column=col_idx)
                # Alternate row shading
                if i % 2 == 0:
                    cell.fill = alternate_fill
                
                # IMPORTANT: Enable wrapping and top-alignment for the large text columns
                if col_idx in [5, 6]: # Summary and Full Body
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                else:
                    cell.alignment = Alignment(vertical="top")

            # Hyperlinks for URL (Column 2)
            url_cell = ws.cell(row=i, column=2)
            url_cell.hyperlink = a.resolved_url or a.url
            url_cell.font = Font(color="0000FF", underline="single")

        # Auto-fit columns with intelligent overrides for content
        for col in range(1, len(headers) + 1):
            column = get_column_letter(col)
            if col in [5, 6]: # Summary and Full Body
                ws.column_dimensions[column].width = 80 # Much wider for content
            else:
                max_length = 0
                for cell in ws[column]:
                    try:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                    except: pass
                adjusted_width = min(max(max_length + 2, 15), 60)
                ws.column_dimensions[column].width = adjusted_width

        # Stream back
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        filename = f"NEXUS_{brand_name}_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
        return StreamingResponse(
            io.BytesIO(output.read()),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )

@router.get("/export/docx")
async def export_docx(
    job_id: str,
    current_user: TokenData = Depends(get_current_user)
):
    """Word document export with full article content."""
    if not HAS_DOCX:
        raise HTTPException(
            status_code=400, 
            detail="The 'python-docx' library is not installed on the server. Please run 'pip install python-docx' to enable Word exports."
        )
    
    async with get_db() as db:
        # Get job for naming
        stmt_job = select(ScrapeJob).where(ScrapeJob.id == job_id)
        if not current_user.is_admin:
            stmt_job = stmt_job.where(ScrapeJob.user_id == current_user.id)
            
        job_res = await db.execute(stmt_job)
        job = job_res.scalar_one_or_none()
        if not job:
            raise HTTPException(404, "Job not found or access denied")
        
        # Get articles
        stmt_articles = select(Article).where(Article.scrape_job_id == job_id)
        if not current_user.is_admin:
            stmt_articles = stmt_articles.where(Article.user_id == current_user.id)
            
        stmt_articles = stmt_articles.order_by(Article.published_at.desc())
        res = await db.execute(stmt_articles)
        articles = res.scalars().all()

        doc = Document()
        
        # Header Section
        title = doc.add_heading(f"NEXUS News Report: {job.sector}", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        metadata = doc.add_paragraph()
        metadata.add_run(f"Report Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        metadata.add_run(f"Sector: {job.sector} | Region: {job.region.upper()}\n")
        metadata.add_run(f"Articles Found: {len(articles)}")
        metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_page_break()

        for a in articles:
            # Article Title
            h = doc.add_heading(a.title, level=1)
            
            # Article Metadata
            p = doc.add_paragraph()
            p.add_run(f"Source: {a.agency or 'Unknown'} | Date: {a.published_at.strftime('%Y-%m-%d') if a.published_at else 'N/A'}\n").italic = True
            p.add_run(f"URL: {a.resolved_url or a.url}").font.color.rgb = RGBColor(0, 0, 255)
            
            # Summary Section
            if a.summary:
                doc.add_heading("Summary", level=2)
                doc.add_paragraph(a.summary)
            
            # Full Content Section
            if a.full_body:
                doc.add_heading("Full Article Content", level=2)
                doc.add_paragraph(a.full_body)
            
            doc.add_paragraph("_" * 50) # Separator

        # Stream back
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        
        brand_name = job.sector.replace(" ", "_")
        filename = f"NEXUS_Report_{brand_name}_{datetime.now().strftime('%Y-%m-%d')}.docx"
        
        return StreamingResponse(
            io.BytesIO(output.read()),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )

@router.get("/{article_id}")
async def get_article(article_id: int, current_user: TokenData = Depends(get_current_user)):
    async with get_db() as db:
        stmt = select(Article).where(Article.id == article_id)
        if not current_user.is_admin:
            stmt = stmt.where(Article.user_id == current_user.id)
            
        res = await db.execute(stmt)
        art = res.scalar_one_or_none()
        if not art: raise HTTPException(404, "Article not found or access denied")
        return art

@router.delete("/bulk")
async def delete_bulk_articles(
    sector: Optional[str] = None,
    region: Optional[str] = None,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    job_id: Optional[str] = None,
    search: Optional[str] = None,
    current_user: TokenData = Depends(get_current_user)
):
    async with get_db() as db:
        stmt = delete(Article)
        if not current_user.is_admin:
            stmt = stmt.where(Article.user_id == current_user.id)
        
        if sector: stmt = stmt.where(Article.sector == sector)
        if region: stmt = stmt.where(Article.region == region)
        if date_from: stmt = stmt.where(Article.published_at >= date_from)
        if date_to: stmt = stmt.where(Article.published_at <= date_to)
        if job_id: stmt = stmt.where(Article.scrape_job_id == job_id)
        if search:
            stmt = stmt.where(or_(
                Article.title.ilike(f"%{search}%"),
                Article.full_body.ilike(f"%{search}%")
            ))
            
        await db.execute(stmt)
        await db.commit()
        return {"status": "success", "message": "Bulk deletion complete"}

@router.delete("/{article_id}")
async def delete_article(article_id: int, current_user: TokenData = Depends(get_current_user)):
    async with get_db() as db:
        stmt = select(Article).where(Article.id == article_id)
        if not current_user.is_admin:
            stmt = stmt.where(Article.user_id == current_user.id)
            
        res = await db.execute(stmt)
        art = res.scalar_one_or_none()
        if not art: raise HTTPException(404, "Article not found or access denied")
        
        await db.execute(delete(Article).where(Article.id == article_id))
        await db.commit()
        return {"status": "success", "message": "Article deleted permanently"}
@router.websocket("/ws/stats")
async def websocket_stats(websocket: WebSocket, token: Optional[str] = Query(None)):
    await websocket.accept()
    if not token:
        await websocket.close(code=status.WS_1008_POLICY_VIOLATION)
        return
    
    from .auth_utils import get_current_user
    try:
        user_data = await get_current_user(token)
    except Exception:
        await websocket.close(code=status.WS_1008_POLICY_VIOLATION)
        return

    try:
        # Send initial stats immediately
        initial_stats = await _fetch_stats_logic(user_data.user_id, user_data.is_admin)
        await websocket.send_json(initial_stats)
        
        while True:
            await asyncio.sleep(10)
            stats = await _fetch_stats_logic(user_data.user_id, user_data.is_admin)
            await websocket.send_json(stats)
    except WebSocketDisconnect:
        pass
    except Exception as e:
        # Ensure we don't crash the worker thread
        print(f"WS Error: {e}")
        pass
